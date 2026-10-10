# -*- coding: utf-8 -*-
"""按 Nature Communications 官方示例格式生成 v32 主稿 Source Data。

示例文件：NC投稿/ncomms-example-source-data.xlsx
  - 每个图面板一个 sheet，命名如 "Figure 1a" / "Figure 5d"
  - sheet 内首行为面板标识，随后为数据表（多数据块用 Parameter/Data 标注）
  - 凝胶/印迹图直接把未裁剪原图贴入 sheet 并标注面板与抗体（本文无此类图）
本文全部为数值图，故每个 sheet = 面板标签 + 表头 + 数值数据。
"""
from __future__ import annotations
import os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, 'source_data', 'source_data_v32.xlsx')
HDR_FILL = PatternFill('solid', fgColor='E8EEF4')
HDR_FONT = Font(bold=True, size=9)
LBL_FONT = Font(bold=True, size=10)


def read(rel):
    return pd.read_csv(os.path.join(ROOT, rel))


def add_sheet(wb, name, title, df, notes=None):
    ws = wb.create_sheet(name[:31])
    ws['A1'] = f'{name}  {title}'
    ws['A1'].font = LBL_FONT
    r0 = 3
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=r0, column=j, value=str(col))
        c.font = HDR_FONT; c.fill = HDR_FILL
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    for i, (_, row) in enumerate(df.iterrows(), start=r0 + 1):
        for j, v in enumerate(row, start=1):
            if isinstance(v, float):
                v = round(v, 6)
            if isinstance(v, (bool, np.bool_)):
                v = 'yes' if v else 'no'
            ws.cell(row=i, column=j, value=v)
    for j, col in enumerate(df.columns, start=1):
        w = max(len(str(col)) * 1.15, 10)
        try:
            w = max(w, min(32, df[col].astype(str).str.len().max() * 1.1))
        except Exception:
            pass
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f'A{r0+1}'
    if notes:
        rr = r0 + len(df) + 2
        for k, v in notes.items():
            ws.cell(row=rr, column=1, value=k).font = Font(bold=True, size=8.5)
            ws.cell(row=rr, column=2, value=v).font = Font(size=8.5)
            rr += 1
    return ws


wb = Workbook(); wb.remove(wb.active)

# ---------- Figure 1 ----------
inter = read('output/experiments/interstudy_comparison/interstudy_60min_comparison.csv')
f1a = inter[['iso3', 'this_arterial_only', 'this_corrected']].copy()
f1a.columns = ['iso3', 'arterial_only_pct', 'full_network_pct']
f1a['abs_difference_pp'] = (f1a.full_network_pct - f1a.arterial_only_pct).abs()
f1a = f1a.sort_values('abs_difference_pp', ascending=False)
add_sheet(wb, 'Figure 1a', 'Same-definition contrast: arterial-only vs with local roads (60 min)', f1a,
          {'median abs difference': '18.76 pp', 'maximum': '75.29 pp (Egypt)'})
f1b = inter[['iso3', 'weiss_2015_cities', 'map_2020_health', 'this_arterial_only', 'this_corrected', 'max_minus_min_pp']].copy()
f1b.columns = ['iso3', 'weiss2015_cities_pct', 'map2020_health_pct', 'this_arterial_pct',
               'this_full_pct', 'range_all_four_pp']
f1b = f1b.sort_values('range_all_four_pp', ascending=False)
add_sheet(wb, 'Figure 1b', 'Cross-product non-comparability (four conventions, 60 min)', f1b,
          {'median range': '30.66 pp', 'maximum': '77.78 pp (EGY)', 'countries >20 pp': '24/30'})

# ---------- Figure 2 / 4：nested decomposition（国家×服务均值） ----------
d = read('output/experiments/direction_law_v2/nested_decomposition.csv')
by = d.groupby('iso3')
agg = pd.DataFrame({
    'nested_arterial_baseline_pct': by['share_60min_pct_A_arterial_nested'].mean(),
    'legacy_arterial_baseline_pct': by['legacy_baseline'].mean(),
    'baseline_difference_pp': by['baseline_diff'].mean(),
    'arterial_coverage_pct': by['arterial_coverage'].mean(),
    'delta_edge_pp': by['delta_edges'].mean(),
    'delta_conv_pp': by['delta_snapping'].mean(),
})
agg['same_baseline_total_pp'] = agg.delta_edge_pp + agg.delta_conv_pp
agg['legacy_mixed_diff_pp'] = by['delta_published'].mean()
agg = agg.reset_index()

f2a = agg[['iso3', 'nested_arterial_baseline_pct', 'legacy_arterial_baseline_pct', 'baseline_difference_pp']]
add_sheet(wb, 'Figure 2a', 'Two arterial-only baselines by country (service-mean)', f2a,
          {'mean baseline difference': '12.13 pp', 'maximum': '72.17 pp (EGY)'})
f2b = agg[['iso3', 'arterial_coverage_pct', 'baseline_difference_pp']].sort_values('baseline_difference_pp', ascending=False)
add_sheet(wb, 'Figure 2b', 'Baseline difference vs arterial coverage', f2b,
          {'Pearson r': '-0.786'})

# ---------- Figure 3：v27 e2_refined（行级 60 行） ----------
e2 = read('output/experiments/v27/e2_refined.csv')
f3a = e2[['iso', 'svc', 'edge', 'snap', 'full']].copy()
f3a.columns = ['iso3', 'service', 'delta_edge_pp', 'delta_conv_pp', 'delta_total_pp']
f3a['abs_delta_conv_pp'] = f3a.delta_conv_pp.abs()
f3a['sign_reverses'] = f3a.abs_delta_conv_pp > f3a.delta_edge_pp
add_sheet(wb, 'Figure 3a', 'Sign decision: |delta_conv| vs delta_edge with 45-degree boundary', f3a,
          {'sign reverses': '31/60 rows', 'sign holds': '29/60 rows'})
f3b = e2[['iso', 'svc', 'snap_loss_pct', 'snap']].copy()
f3b.columns = ['iso3', 'service', 'capture_loss_rate_pct', 'delta_conv_pp']
add_sheet(wb, 'Figure 3b', 'Capture-loss rate vs delta_conv (collinear)', f3b,
          {'Pearson r': '-0.976'})
f3c = e2[['iso', 'svc', 'A', 'full']].copy()
f3c.columns = ['iso3', 'service', 'arterial_baseline_pct', 'delta_total_pp']
f3c['group'] = np.where(f3c.delta_total_pp < 0, 'sign reverses', 'sign holds')
summ = f3c.groupby('group')['arterial_baseline_pct'].agg(['count', 'mean']).reset_index()
summ.columns = ['group', 'n_rows', 'mean_arterial_baseline_pct']
add_sheet(wb, 'Figure 3c', 'Arterial baseline reach by sign group', summ,
          {'sign reverses mean': '93.7%', 'sign holds mean': '76.4%'})

# ---------- Figure 4 ----------
f4a = agg[['iso3', 'delta_edge_pp', 'delta_conv_pp']].sort_values('delta_edge_pp', ascending=False)
add_sheet(wb, 'Figure 4a', 'Edge and capture components on a common baseline (service-mean)', f4a,
          {'delta_edge': '>=0 in 60/60 service rows', 'delta_conv': '<=0 in 60/60 service rows'})
f4b = agg[['iso3', 'same_baseline_total_pp', 'legacy_mixed_diff_pp']].sort_values('same_baseline_total_pp')
add_sheet(wb, 'Figure 4b', 'Total contrast: same baseline vs legacy mixed baseline', f4b,
          {'negative countries (>=1 service)': '19/30', 'negative service rows': '31/60'})
f4c = agg[['iso3', 'arterial_coverage_pct', 'legacy_mixed_diff_pp', 'same_baseline_total_pp']]
add_sheet(wb, 'Figure 4c', 'Coverage relationship: legacy R2=0.77 vs common-baseline R2=0.06', f4c,
          {'legacy OLS R2 (country-mean)': '0.77', 'same-baseline OLS R2': '0.06'})

# ---------- Figure 5 ----------
reg = read('output/experiments/facility_registry/facility_registry_bottleneck.csv')
f5a = reg[['iso3', 'service', 'registry_gain_60min_pp']].sort_values('registry_gain_60min_pp', ascending=False)
add_sheet(wb, 'Figure 5a', 'Independent facility-registry gains (12 cases)', f5a)
env = read('output/experiments/uncertainty/uncertainty_envelope_60min.csv')
env['abs_len'] = env.length_effect_pp.abs()
f5b = env.sort_values('abs_len', ascending=False).head(15)[
    ['iso3', 'service', 'length_effect_pp', 'pop_effect_pp']]
add_sheet(wb, 'Figure 5b', 'Measurement conventions: length metric and population units (top 15)', f5b)
fq = read('output/experiments/experiment_29_global_expansion/facility_quality_by_country.csv')
f5c = fq[fq.raw_records > 0][['country', 'service', 'usable_of_unique_pct']]
f5c.columns = ['country', 'service', 'usable_of_unique_pct']
add_sheet(wb, 'Figure 5c', 'Facility-tag completeness distribution', f5c)
sc = read('output/experiments/facility_scenarios/facility_scenario_shares.csv')
p = sc[sc.service == 'health'].pivot_table(index='iso3', columns='scenario', values='share_60min_pct')
p['hospital_only_penalty_pp'] = p.hospital_clinic - p.hospital_only
f5d = p.reset_index()[['iso3', 'hospital_only_penalty_pp']].sort_values('hospital_only_penalty_pp', ascending=False)
add_sheet(wb, 'Figure 5d', 'Facility-level definition: hospital-only penalty (health, top 14)', f5d)

# ---------- Figure 6 ----------
fl = read('explorations/output/d09_access_identity_flip_summary.csv')
e60 = fl[(fl.service == 'either') & (fl.threshold_min == 60)]
f6a = e60[['iso3', 'total_population', 'flip_population']].copy()
f6a['flip_share_pct'] = 100 * f6a.flip_population / f6a.total_population
f6a = f6a.sort_values('flip_population', ascending=False)
add_sheet(wb, 'Figure 6a', 'Population whose 60-min access status depends on road-network convention', f6a,
          {'total flip population': '984,982,000 (28.6% of 3.45 billion)'})
conn = {}
for _, r in read('explorations/output/d04_connectivity_all.csv').iterrows():
    if float(r.get('bridge_length_share', -1)) >= 0:
        conn[r.iso3] = float(r.connectivity_share)
sub = f6a[f6a.iso3.isin(conn)].copy()
sub['largest_cc_node_share_pct'] = sub.iso3.map(conn) * 100
f6b = sub[['iso3', 'largest_cc_node_share_pct', 'flip_share_pct']]
add_sheet(wb, 'Figure 6b', 'Flip share vs largest-connected-component node share (n=13)', f6b,
          {'Pearson r': '-0.33'})

# ---------- Table 2：外部靶子 ----------
e4 = read('output/experiments/v31/e4_external_target.csv')
t1 = e4[['pair', 'slope', 'r2', 'pearson', 'spearman', 'loco_r2', 'loco_dir', 'n_neg']].copy()
t1.columns = ['difference_form', 'slope', 'OLS_R2', 'pearson_r', 'spearman_rho',
              'LOCO_R2', 'direction_accuracy_pct', 'n_negative_of_30']
t1['difference_form'] = t1.difference_form.str.replace(r'\s*\(.*\)', '', regex=True).str.strip()
add_sheet(wb, 'Table 2', 'External-target test: five difference forms (30 countries)', t1,
          {'source': 'output/experiments/v31/e4_external_target.csv'})

# ---------- Table 1：§2.2 estimand 汇总表 ----------
rows = []
for label, col in [('delta_edge (pure edge increment)', 'delta_edges'),
                   ('delta_conv (capture rule)', 'delta_snapping'),
                   ('delta_total (common baseline)', None),
                   ('delta_pub (mixed baseline)', 'delta_published')]:
    s = d[col] if col else d.delta_edges + d.delta_snapping
    rows.append({'estimand': label, 'min_pp': s.min(), 'median_pp': s.median(),
                 'mean_pp': s.mean(), 'max_pp': s.max(),
                 'negative_rows_of_60': int((s < 0).sum())})
t0 = pd.DataFrame(rows)
add_sheet(wb, 'Table 1', 'Decomposition summary statistics (Section 2.2)', t0,
          {'identity': 'delta_pub = delta_edge + delta_conv + R holds to numerical precision (max dev 7.1e-15)'})

wb.move_sheet('Table 1', offset=-1)
wb.save(OUT)
print('wrote', OUT, '%.1f KB' % (os.path.getsize(OUT) / 1024))
print('sheets:', wb.sheetnames)

