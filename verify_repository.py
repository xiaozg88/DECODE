# -*- coding: utf-8 -*-
"""Integrity and claim checks for the public DECODE GitHub release."""
from __future__ import annotations
import csv, hashlib, os, re
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
MAN=ROOT/'MANIFEST.csv'
errors=[]
checks=[]
def ok(name,cond,detail=''):
 checks.append((name,bool(cond),detail))
 if not cond: errors.append(f'{name}: {detail}')

# Manifest existence and SHA-256.
with open(MAN,newline='',encoding='utf-8') as f:
 rows=list(csv.DictReader(f))
ok('manifest_nonempty',len(rows)>100,f'{len(rows)} entries')
for r in rows:
 p=ROOT/r['file']
 if not p.exists():
  errors.append(f"missing {r['file']}"); continue
 h=hashlib.sha256(p.read_bytes()).hexdigest()
 if h!=r['sha256']: errors.append(f"hash mismatch {r['file']}")

# Required code/data/figures.
code=list(ROOT.glob('*.py'))+list((ROOT/'explorations/scripts').glob('*.py'))
ok('code_count',len(code)>=93,f'{len(code)} Python files')
ok('source_data',(ROOT/'source_data/source_data_v32.xlsx').exists())
ok('main_figures',len(list((ROOT/'figures/main').glob('*.png')))==6 and len(list((ROOT/'figures/main').glob('*.pdf')))==6)
ok('nested_outputs',len(list((ROOT/'output/experiments/nested_counterfactual').glob('*_nested_counterfactual.csv')))==30)

# No hardcoded private workspace path in distributable Python.
abs_path=[]
for p in code:
 try: s=p.read_text(encoding='utf-8-sig')
 except UnicodeDecodeError: continue
 if re.search(r'[A-Za-z]:[\\/](?:nature|Users|ProgramData)',s): abs_path.append(str(p))
ok('no_private_absolute_paths',not abs_path,', '.join(abs_path[:5]))

# Core numerical claims.
d=pd.read_csv(ROOT/'output/experiments/direction_law_v2/nested_decomposition.csv')
ok('nested_60_rows',len(d)==60,f'n={len(d)}')
ok('delta_edge_min',round(float(d.delta_edges.min()),2)==0.22,f"{d.delta_edges.min()}")
ok('delta_edge_negative',int((d.delta_edges<0).sum())==0)
ok('delta_conv_negative',int((d.delta_snapping<0).sum())==60)
delta_total=d.delta_edges+d.delta_snapping
ok('delta_total_negative',int(delta_total.lt(0).sum())==31)
i=pd.read_csv(ROOT/'output/experiments/interstudy_comparison/interstudy_60min_comparison.csv')
ok('interstudy_countries',len(i)==30,f'n={len(i)}')
ok('interstudy_median',round(float(i.max_minus_min_pp.median()),1)==30.7,f'{i.max_minus_min_pp.median():.2f}')
ok('interstudy_max',round(float(i.max_minus_min_pp.max()),1)==77.8)
e=pd.read_csv(ROOT/'output/experiments/v31/e4_external_target.csv')
m=e[e.pair.str.contains('MAP2020') & e.pair.str.contains('主干')].iloc[0]
ok('external_r2',round(float(m.r2),3)==0.754)
ok('external_loco',round(float(m.loco_r2),3)==0.717)

# Source Data structure.
wb=load_workbook(ROOT/'source_data/source_data_v32.xlsx',read_only=True,data_only=True)
ok('source_data_sheets',len(wb.sheetnames)==18,f'{len(wb.sheetnames)}')
ok('source_data_table_order',wb.sheetnames[-2:]==['Table 1','Table 2'],str(wb.sheetnames[-2:]))

report=['DECODE GITHUB RELEASE AUDIT','='*60]
for name,passed,detail in checks: report.append(f"[{'PASS' if passed else 'FAIL'}] {name} {detail}")
report += ['='*60, f'Checks: {sum(x[1] for x in checks)}/{len(checks)} passed']
if errors:
 report += ['FAILURES:']+errors
 print('\n'.join(report)); raise SystemExit(1)
report += ['RESULT: PASS']
print('\n'.join(report))
